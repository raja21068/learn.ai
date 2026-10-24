import pathlib, json, zipfile, subprocess, sys, textwrap, collections
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
ROOT=pathlib.Path('/mnt/data/ai-learning-center')
d=json.loads((ROOT/'assets/catalog.json').read_text())

def rgb(v):
    v=v.lstrip('#');return RGBColor(*[int(v[i:i+2],16) for i in (0,2,4)])
INK=rgb('#202b46');PALE=rgb('#f6f7fb');MUTED=rgb('#6d7893');WHITE=RGBColor(255,255,255)

def textbox(slide,x,y,w,h,txt,size=20,bold=False,color=INK,align=None,font='Aptos'):
    shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=Inches(.06);tf.margin_top=Inches(.04)
    paragraphs=str(txt).split('\n')
    for i,line in enumerate(paragraphs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.space_after=Pt(8)
        if align:p.alignment=align
        for run in p.runs:run.font.name=font;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color
    return shape

def rounded(slide,x,y,w,h,color,border=None,r=0.15):
    box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid();box.fill.fore_color.rgb=rgb(color) if isinstance(color,str) else color
    if border:box.line.color.rgb=rgb(border) if isinstance(border,str) else border
    else:box.line.fill.background()
    return box

def base(prs,accent,number):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=PALE
    rounded(slide,.56,.46,.16,.16,accent)
    textbox(slide,.83,.35,5,.4,'FREE AI LEARNING CENTER  /  QUICK START',11,True,MUTED)
    textbox(slide,12.0,.35,.72,.4,str(number).zfill(2),11,True,MUTED,PP_ALIGN.RIGHT)
    return slide

for idx,c in enumerate(d['courses'],1):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    a=c['accent']
    # Slide 1
    s=base(prs,a,1);rounded(s,.8,1.25,2.05,.42,a);textbox(s,.94,1.28,1.85,.3,f'PATH {idx:02d} / BEGINNER',11,True,WHITE)
    textbox(s,.82,2.2,10.5,1.25,c['title'],52,True,INK)
    textbox(s,.83,3.69,9.2,1.0,c['subtitle'],24,False,MUTED)
    rounded(s,.84,5.6,6.35,.75,'#e9edfa');textbox(s,1.04,5.8,6,.34,'4 MINI LESSONS      ·      20 MINUTES      ·      REAL CODE',14,True,INK)
    # Slide 2
    s=base(prs,a,2);textbox(s,.84,1.13,11,1,'Your four big ideas',39,True,INK)
    for j,l in enumerate(c['lessons']):
        x=.84+(j%2)*6.15;y=2.5+(j//2)*1.75
        rounded(s,x,y,5.75,1.42,WHITE,border='#e4e9f0')
        textbox(s,x+.23,y+.15,.5,.38,str(j+1).zfill(2),16,True,rgb(a))
        textbox(s,x+.88,y+.12,4.58,.4,l['title'],19,True,INK)
        textbox(s,x+.89,y+.59,4.55,.65,l['goal'],12,False,MUTED)
    # Slide 3
    s=base(prs,a,3);textbox(s,.84,1.14,11,1,'What you should remember',38,True,INK)
    pts=[l['takeaway'] for l in c['lessons']]
    for j,p in enumerate(pts):
        y=2.05+j*1.06;rounded(s,.86,y,.49,.49,a);textbox(s,1.5,y-.01,10.75,.75,p,16,False,INK)
    # Slide 4 code
    s=base(prs,a,4);textbox(s,.84,1.1,11,1,'A tiny working example',38,True,INK)
    rounded(s,.85,2.2,11.62,4.25,'#202b46')
    code=c['lessons'][0]['code'].strip();lines=code.splitlines();
    # limit code to available space without hiding all
    textbox(s,1.15,2.6,10.9,3.5,code,size=19 if max(map(len,lines))<65 else 16,color=rgb('#dfecfb'),font='Consolas')
    # slide 5
    s=base(prs,a,5);textbox(s,.84,1.1,11,1,'Your five-minute challenge',37,True,INK)
    l=c['lessons'][0]
    textbox(s,.86,2.09,11.55,1.0,l['challenge']['question'],22,True,INK)
    for j,opt in enumerate(l['challenge']['options']):
        y=3.15+j*1.0;rounded(s,.9,y,11.55,.78,WHITE,border='#e0e5ef');textbox(s,1.13,y+.13,10.8,.5,chr(65+j)+'.  '+opt,18,False,MUTED)
    # slide 6
    s=base(prs,a,6);textbox(s,.84,1.08,11,1,'Check your thinking',38,True,INK)
    l=c['lessons'][0];correct=l['challenge']['correct']
    rounded(s,.86,2.3,11.58,1.18,'#dff7e4');textbox(s,1.12,2.62,11,.7,chr(65+correct)+'.  '+l['challenge']['options'][correct],23,True,rgb('#287146'))
    textbox(s,.94,3.84,11.4,1.65,l['challenge']['explanation'],20,False,INK)
    textbox(s,.94,5.85,10.8,.55,'Next: open index.html and complete the interactive lesson.',16,True,rgb(a))
    prs.save(ROOT/'downloads/slides'/f'{c["id"]}-quickstart.pptx')

for c in d['courses']:
    for l in c['lessons']:
        stem=c['id']+'-'+l['id']
        (ROOT/'downloads/code'/f'{stem}.py').write_text(f'# Free AI Learning Center | {c["title"]} | {l["title"]}\n# MIT licensed original example. Python 3, standard library only.\n\n'+l['code'],encoding='utf-8')
        challenge=l['challenge'];answer='ABC'[challenge['correct']]
        cells=[
         dict(cell_type='markdown',metadata={},source=[f'# {c["title"]}: {l["title"]}\n',f'**Time:** {l["minutes"]} minutes  \n',f'**Objective:** {l["goal"]}\n','\n',l['explain']+'\n','\n',f'> **Analogy:** {l["analogy"]}\n']),
         dict(cell_type='markdown',metadata={},source=['## Two ideas to keep in mind\n','\n']+[f'- {p}\n' for p in l['points']]),
         dict(cell_type='code',metadata={},execution_count=None,outputs=[],source=l['code'].splitlines(True)),
         dict(cell_type='markdown',metadata={},source=['## Five-minute hands-on exercise\n','\n',l['practice']+'\n','\n','### Quick knowledge check\n','\n',challenge['question']+'\n','\n']+[f'- **{chr(65+j)}.** {v}\n' for j,v in enumerate(challenge['options'])]),
         dict(cell_type='markdown',metadata={},source=['<details>\n',f'<summary>Reveal answer: {answer}</summary>\n','\n',challenge['explanation']+'\n','</details>\n','\n',f'**Takeaway:** {l["takeaway"]}\n']),
        ]
        nb=dict(cells=cells,metadata=dict(kernelspec=dict(display_name='Python 3',language='python',name='python3'),language_info=dict(name='python')),nbformat=4,nbformat_minor=5)
        (ROOT/'downloads/notebooks'/f'{stem}.ipynb').write_text(json.dumps(nb,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')

license_text='''# Licensing and source notes

**Original lesson text and original slide decks** in this downloadable Free AI Learning Center are offered under the [Creative Commons Attribution 4.0 International license](https://creativecommons.org/licenses/by/4.0/). Attribution: Free AI Learning Center (2026).

**Original Python code and notebooks** included here are available under the MIT License below.

Copyright (c) 2026 Free AI Learning Center contributors

Permission is hereby granted, free of charge, to any person obtaining a copy of this software and associated documentation files (the "Software"), to deal in the Software without restriction, including without limitation the rights to use, copy, modify, merge, publish, distribute, sublicense, and/or sell copies of the Software, and to permit persons to whom the Software is furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY, FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM, OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE SOFTWARE.

External GitHub repositories and their slide files are **linked but not embedded or relicensed**. Verify the original owners' rights and license terms before redistributing those files. Stanford CS329Z screenshots supplied as topic references are not included here or claimed as original.
'''
(ROOT/'LICENSE.md').write_text(license_text)
readme='''# Free AI Learning Center

A responsive static learning website containing **5 beginner learning paths, 20 original lessons, 20 downloadable Jupyter notebooks, 20 runnable Python scripts, and five original six-slide PPTX decks**.

## Open the website

Open `index.html` in a modern browser. There is no server, package install, account, or API key required to read lessons and take challenges. You can also start a local server:

```bash
python -m http.server 8000
```

Then open http://localhost:8000/ .

The website uses HTML/CSS/vanilla JS and hash-based routes so deep links work on GitHub Pages. It stores completed challenges in browser `localStorage` only. No analytics or external JavaScript libraries.

## Publish free on GitHub Pages

1. Create a new public GitHub repository such as `ai-learning-center`.
2. Upload the **contents of this directory** to the repository root (not the parent folder).
3. In **Settings → Pages**, select **Deploy from a branch**, choose `main` and `/ (root)`, then save.
4. When GitHub provides a Pages URL, open it and test the lessons and downloads.

If publishing under a repository subpath, relative asset paths and hash routes already support that layout. No build step is required.

## What's included

- `index.html`: entry point and page shell
- `assets/style.css`, `assets/app.js`, `assets/catalog.js`: UI and curriculum
- `downloads/notebooks/`: 20 Jupyter notebooks, one per lesson
- `downloads/code/`: 20 standalone Python 3 scripts
- `downloads/slides/`: five original PPTX decks
- `downloads/starter-pack.zip`: combined teaching resources (no duplicate site files)
- `LICENSE.md`: licensing details and external material caveat

Run every code example with `python downloads/code/python-hello-variables.py` (adjust path for the lesson). All examples only use the standard library; LLM/agent examples are deliberately small *simulations of concepts*, not actual calls to an LLM API.

## Verified public repositories linked in the website

- https://github.com/raja21068/Machine-Learning-For-Beginerrs-Course-PPT
- https://github.com/raja21068/Network-Security-Course_PPT
- https://github.com/raja21068/Big-Data-Analytics-and-Mining-Course-PPT
- https://github.com/raja21068/Federated-Learning-Framework (alternative to unavailable requested PPT repository)
- https://github.com/raja21068/Smart-Contract-Blockchain

**Important:** The originally supplied `raja21068/Federated-learning-PPT` URL returned a 404 during verification. The framework repo is a verified *alternative*, not that missing slide collection. External repository resources remain linked at their original source; their license has not been assumed.

The original agentic-AI lessons draw on the concepts of reasoning, memory, environments, tools, and evaluation. They do not reproduce the supplied Stanford CS329Z screenshots or claim Stanford affiliation.

## Content editing

Edit the data in `assets/catalog.json`, then update `assets/catalog.js` so the browser sees identical data. `build_content.py` is the source authoring script for the present catalog; `generate_downloads.py` generates the original notebooks, scripts and slides from it (requires `python-pptx` for slide generation). The deployed site needs neither script nor Python.
'''
(ROOT/'README.md').write_text(readme)

kit=ROOT/'downloads/starter-pack.zip'
with zipfile.ZipFile(kit,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=7) as z:
    for dir_name in ['code','notebooks','slides']:
        for file in sorted((ROOT/'downloads'/dir_name).glob('*')):
            z.write(file,file.relative_to(ROOT))
    z.write(ROOT/'LICENSE.md','LICENSE.md')

print('Slide decks:',len(list((ROOT/'downloads/slides').glob('*.pptx'))))
print('Notebooks:',len(list((ROOT/'downloads/notebooks').glob('*.ipynb'))))
print('Scripts:',len(list((ROOT/'downloads/code').glob('*.py'))))
print('Kit bytes:',kit.stat().st_size)
# Validate scripts and notebooks
for f in sorted((ROOT/'downloads/code').glob('*.py')):
    r=subprocess.run([sys.executable,str(f)],capture_output=True,text=True,timeout=5)
    if r.returncode:print('SCRIPT FAILED',f.name,r.stderr);raise SystemExit(1)
    print('PASS',f.name,r.stdout.strip().replace('\n',' / ')[:95])
for f in sorted((ROOT/'downloads/notebooks').glob('*.ipynb')):
    obj=json.loads(f.read_text());assert obj['nbformat']==4 and len(obj['cells'])==5
print('All tests passed')
